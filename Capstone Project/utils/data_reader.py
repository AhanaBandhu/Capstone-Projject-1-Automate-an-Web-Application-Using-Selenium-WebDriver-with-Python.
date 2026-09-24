"""
Reads test data from Excel (if present) or JSON (default/fallback).
Returns a plain dict so the rest of the framework doesn't care which
source it came from.
"""

import json
import os
import sys
import time

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from config import config


def _load_from_json() -> dict:
    with open(config.JSON_TEST_DATA, "r", encoding="utf-8") as f:
        return json.load(f)


def _load_from_excel() -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(config.EXCEL_TEST_DATA, data_only=True)
    sheet = wb["TestData"] if "TestData" in wb.sheetnames else wb.active

    headers = [cell.value for cell in sheet[1]]
    values = [cell.value for cell in sheet[2]]
    row = dict(zip(headers, values))

    return {
        "user": {
            "first_name": row.get("first_name"),
            "last_name": row.get("last_name"),
            "email_template": row.get("email_template"),
            "telephone": str(row.get("telephone")),
            "password": row.get("password"),
        },
        "product": {
            "search_term": row.get("search_term"),
            "quantity_to_add": int(row.get("quantity_to_add", 1)),
            "updated_quantity": int(row.get("updated_quantity", 1)),
        },
    }


def get_test_data() -> dict:
    if os.path.exists(config.EXCEL_TEST_DATA):
        data = _load_from_excel()
    else:
        data = _load_from_json()

    # Resolve the {timestamp} placeholder so every run registers a
    # unique, non-colliding account on the demo site.
    timestamp = str(int(time.time()))
    data["user"]["email"] = data["user"]["email_template"].format(timestamp=timestamp)
    return data
